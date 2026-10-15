"""
Synthetic Financial Data Generator for Arctus.
Generates realistic Indian mid-market P&L datasets (Auto Ancillary, Specialty Chemicals, B2B SaaS)
including edge-case spreadsheets with merged cells, multi-sheet structures, accounting parentheses,
and transposed orientations for testing and demo flows.
"""

import io
from typing import Dict, Any, List
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from faker import Faker
from core.models import FinancialStatement

fake = Faker("en_IN")


def generate_clean_midmarket_statement(
    company_name: str = "Apex Precision Automotives Ltd.",
    base_quarterly_rev: float = 125000000.0,  # ₹12.5 Cr / quarter (~₹50 Cr annual)
    num_quarters: int = 8
) -> FinancialStatement:
    """
    Generates an in-memory validated FinancialStatement instance for instant demo flows.
    """
    quarters = [
        "Q1 FY23", "Q2 FY23", "Q3 FY23", "Q4 FY23",
        "Q1 FY24", "Q2 FY24", "Q3 FY24", "Q4 FY24"
    ][:num_quarters]

    # Mild quarterly growth trend with realistic operational variance
    rng = np.random.default_rng(seed=101)
    growth_factors = np.cumprod(1.0 + rng.normal(0.025, 0.015, size=num_quarters))
    revenue = base_quarterly_rev * growth_factors

    cogs_pct = rng.normal(0.58, 0.01, size=num_quarters)
    cogs = revenue * cogs_pct
    gross_profit = revenue - cogs

    opex_pct = rng.normal(0.24, 0.008, size=num_quarters)
    opex = revenue * opex_pct
    ebitda = gross_profit - opex

    deprec_interest = ebitda * rng.normal(0.22, 0.01, size=num_quarters)
    pbt = ebitda - deprec_interest
    pat = pbt * 0.75  # 25% tax

    return FinancialStatement(
        periods=quarters,
        revenue=revenue,
        cogs=cogs,
        gross_profit=gross_profit,
        opex=opex,
        ebitda=ebitda,
        pbt=pbt,
        pat=pat,
    )


def create_messy_multisheet_excel(
    company_name: str = "Bharat Precision Engineering Pvt Ltd"
) -> bytes:
    """
    Generates a realistic multi-sheet Excel file with:
    - Sheet 1: Executive Cover Page with merged cells and decorative banners
    - Sheet 2: Unaudited Standalone P&L with Indian currency formats, negative parentheses, and merged headers
    - Sheet 3: Balance Sheet Summary
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Cover
    ws_cover = wb.active
    ws_cover.title = "Company Information"
    ws_cover.views.sheetView[0].showGridLines = True

    ws_cover.merge_cells("B2:F3")
    cover_title = ws_cover["B2"]
    cover_title.value = f"{company_name} — Financial Records"
    cover_title.font = Font(name="Arial", size=14, bold=True, color="14213D")
    cover_title.alignment = Alignment(horizontal="center", vertical="center")

    ws_cover["B5"] = "CIN:"
    ws_cover["C5"] = f"U{fake.random_number(digits=5)}MH2014PTC{fake.random_number(digits=6)}"
    ws_cover["B6"] = "Registered Office:"
    ws_cover["C6"] = "MIDC Industrial Area, Pune, Maharashtra 411018"
    ws_cover["B7"] = "Reporting Currency:"
    ws_cover["C7"] = "INR (₹ in Actuals)"
    ws_cover["B8"] = "Auditors:"
    ws_cover["C8"] = "K. S. Ramanathan & Co., Chartered Accountants"

    # Sheet 2: Standalone P&L Statement (The target sheet)
    ws_pl = wb.create_sheet(title="Statement of Profit & Loss")
    ws_pl.views.sheetView[0].showGridLines = True

    # Decorative header block (to test strip_decorative_headers)
    ws_pl.merge_cells("B2:J2")
    ws_pl["B2"] = company_name.upper()
    ws_pl["B2"].font = Font(name="Arial", size=13, bold=True)
    ws_pl["B2"].alignment = Alignment(horizontal="center")

    ws_pl.merge_cells("B3:J3")
    ws_pl["B3"] = "STATEMENT OF STANDALONE UNAUDITED FINANCIAL RESULTS (₹ in Lakhs)"
    ws_pl["B3"].font = Font(name="Arial", size=10, italic=True)
    ws_pl["B3"].alignment = Alignment(horizontal="center")

    # Table Header at Row 5
    headers = [
        "Particulars", "Q1 FY23", "Q2 FY23", "Q3 FY23", "Q4 FY23",
        "Q1 FY24", "Q2 FY24", "Q3 FY24", "Q4 FY24"
    ]
    for col_idx, h in enumerate(headers, start=2):
        cell = ws_pl.cell(row=5, column=col_idx, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="14213D", end_color="14213D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")

    # Raw line items in Indian accounting labels (values in Lakhs)
    raw_data = [
        ("I. Revenue from Operations", [1150.0, 1220.5, 1280.0, 1350.2, 1310.0, 1390.8, 1450.0, 1520.4]),
        ("II. Other Operating Income", [12.5, 14.0, 15.2, 18.0, 16.5, 17.2, 19.0, 22.1]),
        ("III. Total Income (I + II)", [1162.5, 1234.5, 1295.2, 1368.2, 1326.5, 1408.0, 1469.0, 1542.5]),
        ("Cost of Materials Consumed", [670.0, 710.2, 745.0, 785.4, 760.0, 805.5, 840.0, 880.2]),
        ("Gross Profit", [492.5, 524.3, 550.2, 582.8, 566.5, 602.5, 629.0, 662.3]),
        ("Employee Benefit Expenses", [185.0, 192.5, 198.0, 205.0, 202.0, 210.5, 218.0, 225.0]),
        ("Other Expenses & Overheads", [95.0, 102.0, 105.5, 112.0, 108.0, 115.0, 120.5, 126.8]),
        ("Total Operating Expenses", [280.0, 294.5, 303.5, 317.0, 310.0, 325.5, 338.5, 351.8]),
        ("Operating EBITDA", [212.5, 229.8, 246.7, 265.8, 256.5, 277.0, 290.5, 310.5]),
        ("Finance Costs & Borrowing Charges", [42.0, 44.5, 45.0, 46.2, 45.5, 47.0, 48.2, 49.0]),
        ("Depreciation and Amortisation", [28.0, 28.5, 29.0, 30.0, 30.5, 31.0, 32.0, 33.0]),
        ("Profit Before Tax (PBT)", [142.5, 156.8, 172.7, 189.6, 180.5, 199.0, 210.3, 228.5]),
        ("Tax Expense (Current & Deferred)", [35.6, 39.2, 43.2, 47.4, 45.1, 49.8, 52.6, 57.1]),
        ("Profit After Tax (PAT)", [106.9, 117.6, 129.5, 142.2, 135.4, 149.2, 157.7, 171.4]),
    ]

    for row_idx, (label, values) in enumerate(raw_data, start=6):
        # Label cell
        ws_pl.cell(row=row_idx, column=2, value=label)
        for col_offset, val in enumerate(values):
            col_idx = 3 + col_offset
            # Format some values with accounting string representation, e.g. "₹ 1,150.00"
            if row_idx == 6:
                cell_val = f"₹ {val:,.2f}"
            elif row_idx in (15,):
                # Negative interest formatting test: (42.00)
                cell_val = f"({val:,.2f})"
            else:
                cell_val = val
            ws_pl.cell(row=row_idx, column=col_idx, value=cell_val)

    # Sheet 3: Balance Sheet (Distractor sheet)
    ws_bs = wb.create_sheet(title="Balance Sheet")
    ws_bs["B2"] = "STATEMENT OF ASSETS AND LIABILITIES"
    ws_bs["B4"] = "Property, Plant & Equipment"
    ws_bs["C4"] = 4500.0
    ws_bs["B5"] = "Inventories"
    ws_bs["C5"] = 1250.0
    ws_bs["B6"] = "Trade Receivables"
    ws_bs["C6"] = 1820.0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_transposed_csv() -> bytes:
    """
    Generates a transposed CSV with periods in the first column and line items as headers.
    """
    data = {
        "Quarter": ["Q1 FY24", "Q2 FY24", "Q3 FY24", "Q4 FY24"],
        "Revenue from Operations": ["₹ 1,250.00 Lakhs", "₹ 1,320.50 Lakhs", "₹ 1,390.00 Lakhs", "₹ 1,480.00 Lakhs"],
        "Cost of Materials": [750.0, 792.3, 834.0, 888.0],
        "Gross Profit": [500.0, 528.2, 556.0, 592.0],
        "Operating Expenses": [290.0, 305.0, 320.0, 340.0],
        "Operating EBITDA": [210.0, 223.2, 236.0, 252.0],
        "Profit Before Tax": [150.0, 160.0, 170.0, 182.0],
        "Profit After Tax": [112.5, 120.0, 127.5, 136.5],
    }
    df = pd.DataFrame(data)
    buf = io.BytesIO()
    df.to_csv(buf, index=False)
    return buf.getvalue()
